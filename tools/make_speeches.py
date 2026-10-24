#!/usr/bin/env python3
"""Defence prompters («СУФЛЕР») for every work of every subject → Здача/<N. дисципліна>/<робота>/Суфлер/NN_<назва>_СУФЛЕР.docx.
The report, its extra materials and Здача/README.md are put next to them by tools/build_submission.py (which calls
this module), so each work's folder holds everything for that work.

Works are listed in WORKS (subject, folder, label, title, content file, kind):
* kind "drawio" (ПІС labs) — one prompter per draw.io file, order = tools/drawio_export.py MANIFEST;
* kind "report" (other subjects) — one prompter per defence topic; each speech names the figures/tables to show.
Content file: SPEECHES = [ {id, title, notation, role, opening, walk, traps, defense, [show]} ... ]; for "drawio"
`id` is the draw.io stem (key "drawio" is accepted too). Inline markup: **жирний**.
Every content file also defines DEMO = {minutes, intro, steps: [(крок, що показати, що сказати)], finish} — the
narration script for demonstrating the whole work → <робота>/Суфлер/00_Сценарій_демонстрації.docx.
For "report" works every figure and table caption of the .docx report must be named in some speech's `show` or
`walk` («рис. 2.4», «табл. 1.7–1.27») — the tool aborts otherwise, so no component is left without a prompter.

    .venv/bin/python tools/make_speeches.py              # all works
    .venv/bin/python tools/make_speeches.py tspp cyb     # selected subjects (pis, tspp, cyb, tau, bank)
"""
import importlib.util
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from drawio_export import MANIFEST  # noqa: E402
from uk_plural import n_uk  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
FONT = "Times New Roman"
OUT = ROOT / "Здача"
SUBJ_DIR = {"pis": "1. Проєктування інформаційних систем", "tspp": "2. Технології створення програмних продуктів",
            "tau": "3. Теорія автоматичного управління", "bank": "4. Системи банкової безпеки",
            "cyb": "5. Кібернетична безпека підприємства"}
SUBJECTS = {"pis": "Проєктування інформаційних систем", "tspp": "Технології створення програмних продуктів",
            "cyb": "Кібернетична безпека підприємства", "tau": "Теорія автоматичного управління",
            "bank": "Системи банкової безпеки"}
LAB_TITLE = {1: "Вибір об’єкта дослідження, дерево цілей і WBS", 2: "Моделювання бізнес-процесів у нотації BPMN",
             3: "Функціональне моделювання IDEF0 / IDEF3", 4: "Моделювання DFD і вартісний аналіз (ABC)",
             5: "Об’єктно-орієнтоване моделювання UML", 7: "Концепція архітектури IoT-системи"}
# (subject, folder, label, title, content file, kind, report file)
WORKS = [("pis", f"ЛР{n}", f"ЛР {n}", t, f"lab{n}/speech/content_lab{n}.py", "drawio", f"lab{n}/ЛР{n}_Кравченко_ІР-31.docx")
         for n, t in LAB_TITLE.items()] + [
    ("pis", "КР", "КР", "Курсова робота", "coursework_final/speech/content_kr.py", "report",
     "coursework_final/КР_Кравченко_ІР-31.docx"),
    ("tspp", "ПР1", "ПР 1", "Збір та документування вимог для програмного продукту", "tspp_work/speech/content_pr1.py",
     "report", "tspp_work/pract1/ПР1_Кравченко_ІР-31.docx"),
    ("tspp", "ПР2", "ПР 2", "Архітектура IoT екосистеми", "tspp_work/speech/content_pr2.py", "report",
     "tspp_work/pract2/ПР2_Кравченко_ІР-31.docx"),
    ("tspp", "ПР3", "ПР 3", "Опис існуючих програмних та апаратних рішень", "tspp_work/speech/content_pr3.py",
     "report", "tspp_work/pract3/ПР3_Кравченко_ІР-31.docx"),
    ("tspp", "ПР4", "ПР 4", "Проектування системи", "tspp_work/speech/content_pr4.py", "report",
     "tspp_work/pract4/ПР4_Кравченко_ІР-31.docx"),
    ("tspp", "ЛР1", "ЛР 1", "Опис бази даних", "tspp_work/speech/content_lr1.py", "report",
     "tspp_work/lab1/ЛР1_Кравченко_ІР-31.docx"),
    ("tspp", "ЛР2", "ЛР 2", "Опис програмної реалізації системи", "tspp_work/speech/content_lr2.py", "report",
     "tspp_work/lab2/ЛР2_Кравченко_ІР-31.docx"),
    ("tspp", "Тези", "Тези", "Тези доповіді на конференцію НУХТ", "tspp_work/speech/content_theses.py", "report",
     "Технології створення програмних продуктів/НУХТ/Готово/Кравченко_П.docx"),
    ("cyb", "ЛР1", "ЛР 1", "Структура системи управління кібербезпекою підприємства та стандарти",
     "cybersec_work/speech/content_lab1.py", "report", "cybersec_work/lab1/ЛР1_Кравченко_ІР-31.docx"),
] + [("tau", f"ПЗ{n}", f"ПЗ {n}", t, f"tau_work/speech/content_pz{n}.py", "report",
      f"ТАУ/ПЗ{n}_Кравченко_ІР-31_варіант8.docx")
     for n, t in ((1, "Передатні функції елементів САУ"), (2, "Часові характеристики ланок САУ"),
                  (3, "Перетворення структурних схем і передатні функції САУ"),
                  (4, "Логарифмічні частотні характеристики розімкненої системи"), (5, "Оцінка точності АСУ"),
                  (6, "Оцінка стійкості АСУ"))] + [
    ("bank", f"ПР{n}", f"ПР {n}", t, f"bank_work/speech/content_pr{n}.py", "report",
     f"Системи банкової безпеки/Практична {n}/Готово/ПР_{n}_Варіант_08_Кравченко.xlsx")
    for n, t in ((1, "Оцінка ризиків для відділення банку"), (2, "Модель загроз шляху платіжної картки"),
                 (3, "Проєкт мережі відділення: зони, VLAN, адреси"))]


def _runs(par, text, size=None, bold=False, italic=False, color=None):
    for i, chunk in enumerate(re.split(r"\*\*", text)):
        if not chunk:
            continue
        r = par.add_run(chunk)
        r.bold = bold or (i % 2 == 1)
        r.italic = italic
        r.font.name = FONT
        r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)


def _shade(cell, fill):
    tcPr = cell._element.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
    tcPr.append(sh)


def _para(doc, text="", size=13, bold=False, italic=False, align=None, before=0, after=4, indent=None, color=None,
          style=None):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before = Pt(before); pf.space_after = Pt(after); pf.line_spacing = 1.12
    if indent is not None:
        pf.left_indent = Cm(indent)
    if align:
        p.alignment = align
    _runs(p, text, size, bold, italic, color)
    return p


def _heading(doc, text, num):
    p = _para(doc, f"{num}. {text}", size=14, bold=True, before=10, after=4, color="1F3864")
    p.paragraph_format.keep_with_next = True
    # bottom rule
    pPr = p._element.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "1F3864")):
        b.set(qn(f"w:{k}"), v)
    bdr.append(b); pPr.append(bdr)


def build(work, idx, total, sp, out):
    subj, folder, label, title, _content, kind, report = work
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    _para(doc, f"СУФЛЕР ДО ЗАХИСТУ · {label} · {title}", size=10.5, italic=True, color="595959", after=2)
    _para(doc, sp["title"], size=17, bold=True, after=6)

    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [("Дисципліна", SUBJECTS[subj])]
    if kind == "drawio":
        rows += [("Нотація", sp["notation"]), ("Рисунок у звіті", sp.get("_fig", "—")),
                 ("Файл draw.io", f"draw.io/{sp['id']}.drawio (у папці роботи)"),
                 ("Порядок показу", f"{idx} з {total}"), ("Роль у системі", sp["role"])]
    else:
        rows += [("Тема відповіді", sp["notation"]), ("Що показати у звіті", sp.get("show", "—")),
                 ("Звіт", f"{Path(report).name} (у папці роботи)"), ("Порядок показу", f"{idx} з {total}"), ("Роль у роботі", sp["role"])]
    for k, v in rows:
        c0, c1 = t.add_row().cells
        c0.width, c1.width = Cm(4.2), Cm(12.8)
        _shade(c0, "E8EEF7")
        c0.paragraphs[0].text = ""; _runs(c0.paragraphs[0], k, 11.5, bold=True)
        c1.paragraphs[0].text = ""; _runs(c1.paragraphs[0], v, 11.5)

    _heading(doc, "З чого почати відповідь", 1)
    for s in (sp["opening"] if isinstance(sp["opening"], list) else [sp["opening"]]):
        p = _para(doc, s, size=13.5, after=4)
        p.paragraph_format.left_indent = Cm(0.4)

    _heading(doc, "Покрокова екскурсія по схемі" if kind == "drawio" else "Покрокова відповідь", 2)
    for i, step in enumerate(sp["walk"], 1):
        p = _para(doc, "", size=13, after=4, indent=0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        _runs(p, f"{i}. ", 13, bold=True, color="1F3864")
        _runs(p, step, 13)

    _heading(doc, "Підводні камені та типові питання викладача", 3)
    for q, a in sp["traps"]:
        p = _para(doc, "", size=13, before=4, after=1)
        p.paragraph_format.keep_with_next = True
        _runs(p, "Питання: ", 13, bold=True, color="C00000")
        _runs(p, q, 13, bold=True)
        p = _para(doc, "", size=13, after=4, indent=0.6)
        _runs(p, "Відповідь: ", 13, bold=True, color="1F6F3C")
        _runs(p, a, 13)

    _heading(doc, "Що відповісти, якщо придереться", 4)
    for d in sp["defense"]:
        p = _para(doc, "", size=13, after=3, indent=0.6)
        p.paragraph_format.first_line_indent = Cm(-0.5)
        _runs(p, "▸ ", 13, color="1F3864")
        _runs(p, d, 13)

    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def build_demo(work, demo, out):
    """Сценарій демонстрації всієї роботи: що показати і що сказати на кожному кроці."""
    subj, folder, label, title, _content, kind, report = work
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2); sec.top_margin = sec.bottom_margin = Cm(1.8)
    st = doc.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(13)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    _para(doc, f"СЦЕНАРІЙ ДЕМОНСТРАЦІЇ · {label} · {title}", size=10.5, italic=True, color="595959", after=2)
    _para(doc, f"{label}: {title}", size=17, bold=True, after=6)
    t = doc.add_table(rows=0, cols=2)
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    files = ("draw.io/*.drawio у папці роботи (порядок — як у кроках)" if kind == "drawio"
             else f"{Path(report).name} у папці роботи")
    rows = [("Дисципліна", SUBJECTS[subj])] + ([("Звіт", Path(report).name)] if kind == "drawio" else [])
    for k, v in rows + [("Що відкрити заздалегідь", files),
                 ("Тривалість", f"≈ {demo['minutes']} хв ({n_uk(len(demo['steps']), 'крок', 'кроки', 'кроків')})"),
                 ("Як користуватися", "Ліва колонка — що показати на екрані, права — що сказати. Говоріть "
                                      "своїми словами; деталі й відповіді на питання — у суфлерах цієї роботи.")]:
        c0, c1 = t.add_row().cells
        c0.width, c1.width = Cm(4.2), Cm(12.8)
        _shade(c0, "E8EEF7")
        c0.paragraphs[0].text = ""; _runs(c0.paragraphs[0], k, 11.5, bold=True)
        c1.paragraphs[0].text = ""; _runs(c1.paragraphs[0], v, 11.5)
    _heading(doc, "Вступ", 1)
    _para(doc, demo["intro"], size=13.5, after=4).paragraph_format.left_indent = Cm(0.4)
    _heading(doc, "Хід демонстрації", 2)
    steps = doc.add_table(rows=1, cols=3)
    steps.style = "Table Grid"; steps.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = (Cm(4.0), Cm(4.6), Cm(8.4))
    for c, w, h in zip(steps.rows[0].cells, widths, ("Крок", "Що показати", "Що сказати")):
        c.width = w; _shade(c, "1F3864")
        c.paragraphs[0].text = ""; _runs(c.paragraphs[0], h, 11.5, bold=True, color="FFFFFF")
    for i, (name, show, say) in enumerate(demo["steps"], 1):
        cells = steps.add_row().cells
        for c, w in zip(cells, widths):
            c.width = w
        _shade(cells[0], "E8EEF7")
        cells[0].paragraphs[0].text = ""; _runs(cells[0].paragraphs[0], f"{i}. {name}", 11.5, bold=True)
        cells[1].paragraphs[0].text = ""; _runs(cells[1].paragraphs[0], show, 11, color="1F3864")
        cells[2].paragraphs[0].text = ""; _runs(cells[2].paragraphs[0], say, 11.5)
    _heading(doc, "Завершення", 3)
    _para(doc, demo["finish"], size=13.5, after=4).paragraph_format.left_indent = Cm(0.4)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def load_demo(path):
    spec = importlib.util.spec_from_file_location(Path(path).stem + "_demo", ROOT / path)
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
    return getattr(mod, "DEMO", None)


def load_tables(path):
    """TABLE_SPEECHES робіт з draw.io: суфлери до таблиць звіту (компонентів, які не є схемами)."""
    spec = importlib.util.spec_from_file_location(Path(path).stem + "_tables", ROOT / path)
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
    return getattr(mod, "TABLE_SPEECHES", [])


def captions(report, kind):
    """Номери рисунків («Р»)/таблиць («Т») звіту в порядку появи."""
    out = []
    for p in Document(ROOT / report).paragraphs:
        m = re.match(r"^(Рисунок|Рис\.|Таблиця)\s+(\d+\.\d+)\b", p.text.strip())
        if m and (("Р" if m.group(1).startswith("Рис") else "Т") == kind):
            out.append(m.group(2))
    return out


def load(path):
    spec = importlib.util.spec_from_file_location(Path(path).stem, ROOT / path)
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
    out = mod.SPEECHES
    for sp in out:
        sp.setdefault("id", sp.get("drawio"))
    return out


def slug(s):
    return re.sub(r"[^\w\-.]+", "_", s, flags=re.U).strip("_")


def work_dir(work):
    """Папка роботи; без символів, заборонених в іменах файлів Windows (: / \\ ? * " < > |)."""
    subj, folder, _label, title, *_ = work
    name = title if title.startswith(folder) else f"{folder} {title}"
    name = re.sub(r"\s*:\s*", " — ", name.replace(" / ", " і ").replace("/", "-"))
    return OUT / SUBJ_DIR[subj] / re.sub(r'[\\?*"<>|]', "", name)


def refs(text):
    """«рис. 1.2–1.5», «табл. 1.7, 1.8», «рисунки 2.1–2.3» → {"Р1.2", …}."""
    out = set()
    for kind, body in re.findall(r"(рис|табл)[а-яіїє]*\.?\s*((?:\d+\.\d+(?:\s*[–-]\s*(?:\d+\.)?\d+)?"
                                 r"(?:\s*(?:,|і|та)\s*)?)+)", text, re.I):
        k = "Р" if kind.lower() == "рис" else "Т"
        for a, b1, b2 in re.findall(r"(\d+)\.(\d+)(?:\s*[–-]\s*(?:\d+\.)?(\d+))?", body):
            out |= {f"{k}{a}.{x}" for x in range(int(b1), int(b2 or b1) + 1)}
    return out


def uncovered(report, speeches):
    """Рисунки й таблиці звіту, яких не названо в жодному суфлері."""
    caps = set()
    for p in Document(ROOT / report).paragraphs:
        m = re.match(r"^(Рисунок|Рис\.|Таблиця)\s+(\d+\.\d+)\b", p.text.strip())
        if m:
            caps.add(("Р" if m.group(1).startswith("Рис") else "Т") + m.group(2))
    cov = set().union(*[refs(" ".join([sp.get("show", "")] + sp["walk"])) for sp in speeches])
    return sorted(caps - cov)


def main(argv):
    """Будує суфлери й сценарії; повертає {work: [(файл, назва)…]} для покажчика tools/build_submission.py."""
    import shutil
    subjects = [a for a in argv if a in SUBJECTS] or list(SUBJECTS)
    index = {}
    for work in WORKS:
        subj, folder, label, title, content, kind, report = work
        if subj not in subjects:
            continue
        speeches = load(content)
        if kind == "drawio":
            names = [m[1] for m in MANIFEST[int(folder[2:])]]
            got = [s["id"] for s in speeches]
            assert got == names, f"{folder}: speech order/names {got} != draw.io manifest {names}"
            figs = captions(report, "Р")
            assert len(figs) == len(speeches), f"{folder}: {len(figs)} рисунків у звіті ≠ {len(speeches)} draw.io"
            for sp, num in zip(speeches, figs):          # схеми draw.io йдуть у звіті в тому самому порядку
                sp["_fig"] = f"рис. {num}"
            tables = load_tables(content)
            miss = sorted(set(f"Т{t}" for t in captions(report, "Т")) -
                          set().union(set(), *[refs(" ".join([t.get("show", "")] + t["walk"])) for t in tables]))
            assert not miss, f"{subj} {folder}: таблиці без суфлера: {miss}"
        elif report.endswith(".docx"):
            miss = uncovered(report, speeches)
            assert not miss, f"{subj} {folder}: рисунки/таблиці без суфлера: {miss}"
        demo = load_demo(content)
        assert demo and all(demo.get(k) for k in ("minutes", "intro", "steps", "finish")), \
            f"{subj} {folder}: no DEMO narration script in {content}"
        assert all(len(s) == 3 and all(s) for s in demo["steps"]), f"{subj} {folder}: DEMO step must be (крок, показати, сказати)"
        sdir = work_dir(work) / "Суфлер"
        shutil.rmtree(sdir, ignore_errors=True)          # без застарілих суфлерів попередніх версій
        dout = sdir / "00_Сценарій_демонстрації.docx"
        build_demo(work, demo, dout)
        items = index.setdefault(work, [(dout, "Сценарій демонстрації всієї роботи")])
        for i, sp in enumerate(speeches, 1):
            for k in ("id", "title", "notation", "role", "opening", "walk", "traps", "defense"):
                assert sp.get(k), f"{subj} {folder} {sp.get('id')}: empty {k}"
            assert 3 <= len(sp["traps"]) <= 5, f"{subj} {folder} {sp['id']}: need 3–5 traps"
            name = sp["id"] if kind == "drawio" else f"{i:02d}_{slug(sp['id'])}"
            out = sdir / f"{name}_СУФЛЕР.docx"
            build(work, i, len(speeches), sp, out)
            items.append((out, sp["title"]))
        if kind == "drawio":                             # суфлери до таблиць звіту — після схем
            rwork = work[:5] + ("report",) + work[6:]
            for i, sp in enumerate(tables, 1):
                for k in ("id", "title", "notation", "role", "opening", "walk", "traps", "defense"):
                    assert sp.get(k), f"{subj} {folder} {sp.get('id')}: empty {k}"
                assert 3 <= len(sp["traps"]) <= 5, f"{subj} {folder} {sp['id']}: need 3–5 traps"
                out = sdir / f"Т{i:02d}_{slug(sp['id'])}_СУФЛЕР.docx"
                build(rwork, i, len(tables), sp, out)
                items.append((out, sp["title"]))
        n_tab = len(tables) if kind == "drawio" else 0
        print(f"{SUBJECTS[subj]} / {folder}: {len(speeches)}" + (f" + {n_tab} табл." if n_tab else "") + " + сценарій")
    return index


if __name__ == "__main__":
    main(sys.argv[1:])
