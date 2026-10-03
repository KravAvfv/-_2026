#!/usr/bin/env python3
"""Assemble a ПІС laboratory report (.docx) from the official title page + a YAML content spec.

Usage:
    .venv/bin/python tools/build_report.py lab1/report.yaml [--pdf]

Spec format (YAML) — see .claude/skills/academic-report-builder/SKILL.md for the full reference:

    lab_number: 1
    output: lab1/ЛР1_Кравченко_ІР-31.docx
    year: 2026                       # optional, default = current year
    theme: "..."                     # optional, default = first line of My_theme.txt
    blocks:
      - h1: Мета роботи
      - p: Текст абзацу з посиланням на {fig:goal_tree} та {tab:units}.
      - h2: 1.1 Підрозділ
      - bullets: [пункт, пункт]
      - numbered: [пункт, пункт]
      - table: {id: units, caption: Назва, header: [..], rows: [[..]], widths: [3, 6, 7], font: 12}
      - figure: {id: goal_tree, src: lab1/diagrams/goal_tree.mmd, caption: Назва, orientation: auto}
      - pagebreak: true

Numbering: figures → "Рис. <lab>.<k> — <caption>", tables → "Таблиця <lab>.<k> — <caption>".
`{fig:id}` / `{tab:id}` in any text are replaced with "<lab>.<k>". Unknown ids abort the build.
Formatting follows ДСТУ 3008:2015 conventions used by the department: Times New Roman 14 pt,
1.5 line spacing, 1.25 cm first-line indent, justified body; page numbers bottom-centre from page 2.
"""
import argparse, copy, datetime, json, re, subprocess, sys
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "Титульна сторінка для лб.docx"
THEME_FILE = ROOT / "My_theme.txt"
FONT = "Times New Roman"
# Page geometry (cm) — keep in sync with tools/mmd_render.py PAGES
PORTRAIT = dict(w=21.0, h=29.7, left=3.0, right=1.5, top=2.0, bottom=2.0)
LANDSCAPE = dict(w=29.7, h=21.0, left=2.0, right=2.0, top=3.0, bottom=1.5)


# ----------------------------------------------------------------------------- helpers
def set_run_font(run, size=14, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts"); rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def fmt_par(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, indent=1.25, spacing=1.5, before=0, after=0,
            keep_next=False, left=0.0):
    pf = p.paragraph_format
    p.alignment = align
    pf.first_line_indent = Cm(indent) if indent else Cm(0)
    pf.left_indent = Cm(left)
    pf.line_spacing = spacing
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = keep_next
    pf.widow_control = True


INLINE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")


def add_text(p, text, size=14, bold=False, italic=False):
    """Supports **bold** and *italic* inline markup."""
    for part in INLINE.split(text):
        if not part:
            continue
        b, i, t = bold, italic, part
        if part.startswith("**") and part.endswith("**"):
            b, t = True, part[2:-2]
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            i, t = True, part[1:-1]
        set_run_font(p.add_run(t), size, b, i)


def set_section_geometry(section, geo, orient):
    section.orientation = orient
    section.page_width, section.page_height = Cm(geo["w"]), Cm(geo["h"])
    section.left_margin, section.right_margin = Cm(geo["left"]), Cm(geo["right"])
    section.top_margin, section.bottom_margin = Cm(geo["top"]), Cm(geo["bottom"])
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def add_page_number_field(paragraph):
    run = paragraph.add_run()
    set_run_font(run, 12)
    for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
        run._element.append(el)


def repeat_table_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trpr.append(el)


def no_split_row(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true"); trpr.append(el)


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "6"); el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def cell_text(cell, text, size, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    lines = text if isinstance(text, list) else str(text).split("\n")
    for k, line in enumerate(lines):
        p = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
        fmt_par(p, align=align, indent=0, spacing=1.0)
        add_text(p, str(line), size=size, bold=bold)


def style_by_name(doc, name, outline=None):
    """python-docx lookups translate UI names; the template stores "Heading 1" verbatim — search by name."""
    for st in doc.styles:
        if st.name == name:
            return st
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    if outline is not None:
        lvl = OxmlElement("w:outlineLvl"); lvl.set(qn("w:val"), str(outline))
        st.element.get_or_add_pPr().append(lvl)
    return st


def normalize_style_font(st, size=14, bold=False):
    """Replace the style's run properties (template: Arial 20 pt blue) with Times New Roman, black."""
    rpr = st.element.get_or_add_rPr()
    for child in list(rpr):
        rpr.remove(child)
    rf = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(attr), FONT)
    rpr.append(rf)
    if bold:
        rpr.append(OxmlElement("w:b"))
    col = OxmlElement("w:color"); col.set(qn("w:val"), "000000"); rpr.append(col)
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag); el.set(qn("w:val"), str(size * 2)); rpr.append(el)


def fld(run_parent, kind, text=None):
    """Append a complex-field part (begin/instr/separate/end) as its own run to `run_parent` (w:p or w:hyperlink)."""
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr"); rf = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rf.set(qn(attr), FONT)
    rpr.append(rf)
    for tag in ("w:sz", "w:szCs"):
        el = OxmlElement(tag); el.set(qn("w:val"), "28"); rpr.append(el)
    r.append(rpr)
    if kind == "instr":
        el = OxmlElement("w:instrText"); el.set(qn("xml:space"), "preserve"); el.text = text
    elif kind == "text":
        el = OxmlElement("w:t"); el.set(qn("xml:space"), "preserve"); el.text = text
    elif kind == "tab":
        el = OxmlElement("w:tab")
    else:
        el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind)
    r.append(el)
    run_parent.append(r)


# ----------------------------------------------------------------------------- builder
class ReportBuilder:
    def __init__(self, spec, spec_dir, break_tables=frozenset(), toc_pages=None):
        self.break_tables = set(break_tables)   # table numbers that must start on a new page (layout loop)
        self.table_meta = []                    # [(key, keep_together, probe, last_row_text, header_text)]
        self.toc_pages = toc_pages or {}        # bookmark -> page number (from the rendered PDF)
        self.toc_entries = []                   # [(level, text, bookmark)] in document order
        self._bm = 0
        self._break_next = False
        self.spec, self.spec_dir = spec, spec_dir
        self.lab = int(spec["lab_number"])
        self.doc = Document(str(TEMPLATE))
        self.fig_ids, self.tab_ids = {}, {}
        self.orientation = "portrait"
        self._number_objects()
        self._prepare_styles()

    # -- numbering pass (so forward references work)
    def _number_objects(self):
        f = t = 0
        for b in self.spec["blocks"]:
            if "figure" in b:
                f += 1
                if b["figure"].get("id"):
                    self.fig_ids[b["figure"]["id"]] = f"{self.lab}.{f}"
            if "table" in b:
                t += 1
                if b["table"].get("id"):
                    self.tab_ids[b["table"]["id"]] = f"{self.lab}.{t}"

    def resolve(self, text):
        def rep(m):
            kind, key = m.group(1), m.group(2)
            table = self.fig_ids if kind == "fig" else self.tab_ids
            if key not in table:
                sys.exit(f"Unknown reference {{{kind}:{key}}}")
            return table[key]
        return re.sub(r"\{(fig|tab):([\w\-]+)\}", rep, str(text))

    def _prepare_styles(self):
        for lvl in (1, 2, 3):
            normalize_style_font(style_by_name(self.doc, f"Heading {lvl}", outline=lvl - 1), bold=True)
        for lvl in (1, 2):
            normalize_style_font(style_by_name(self.doc, f"toc {lvl}"))

    def toc(self):
        """ЗМІСТ: native TOC field (TOC \\o "1-3" \\h \\z \\u) whose cached result is pre-filled with clickable
        entries (w:hyperlink → heading bookmark) and PAGEREF page numbers taken from the rendered PDF, so it is
        correct on opening in any editor and can be regenerated with "Update field"."""
        head = self.doc.add_paragraph()
        fmt_par(head, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, before=0, after=12)
        add_text(head, "ЗМІСТ", bold=True)
        entries = [(lvl, txt, f"_Toc{k:06d}") for k, (lvl, txt) in enumerate(self._scan_headings(), 1)]
        avail = PORTRAIT["w"] - PORTRAIT["left"] - PORTRAIT["right"]
        for i, (lvl, txt, bm) in enumerate(entries):
            par = self.doc.add_paragraph()
            par.style = style_by_name(self.doc, f"toc {min(lvl, 2)}")
            fmt_par(par, align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, spacing=1.5, left=0 if lvl == 1 else 0.75)
            par.paragraph_format.tab_stops.add_tab_stop(Cm(avail), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            if i == 0:
                fld(par._p, "begin"); fld(par._p, "instr", ' TOC \\o "1-3" \\h \\z \\u '); fld(par._p, "separate")
            link = OxmlElement("w:hyperlink"); link.set(qn("w:anchor"), bm); link.set(qn("w:history"), "1")
            fld(link, "text", txt); fld(link, "tab")
            fld(link, "begin"); fld(link, "instr", f" PAGEREF {bm} \\h "); fld(link, "separate")
            fld(link, "text", str(self.toc_pages.get(bm, ""))); fld(link, "end")
            par._p.append(link)
            if i == len(entries) - 1:
                fld(par._p, "end")
        self._break_next = True                  # body starts on the next page

    def _scan_headings(self):
        out = []
        for b in self.spec["blocks"]:
            key = next(iter(b))
            if key in ("h1", "h2"):
                txt = self.number_dot(self.resolve(b[key]))
                out.append((1 if key == "h1" else 2, txt.upper() if key == "h1" else txt))
        return out

    # -- title page
    def fill_title(self):
        theme = self.spec.get("theme") or THEME_FILE.read_text(encoding="utf-8").strip().splitlines()[0]
        year = self.spec.get("year") or datetime.date.today().year
        ps = self.doc.paragraphs
        done = {"lab": False, "theme": False, "year": False}
        for p in ps:
            txt = p.text
            if txt.startswith("Лабораторна робота №") and not done["lab"]:
                for r in p.runs:
                    if r.text.strip("_ ") == "" and "_" in r.text:
                        r.text = str(self.lab); r.bold = True; done["lab"] = True; break
            elif txt.startswith("Тема:") and not done["theme"]:
                for r in p.runs:
                    if r.text.strip("_ ") == "" and "_" in r.text:
                        r.text = theme; done["theme"] = True; break
            elif re.fullmatch(r"\s*Київ-\d{4}\s*", txt) and not done["year"]:
                for r in p.runs:
                    if "Київ" in r.text:
                        r.text = f"Київ-{year}"; done["year"] = True
        missing = [k for k, v in done.items() if not v]
        if missing:
            sys.exit(f"Title page placeholders not found: {missing}")
        cp = self.doc.core_properties
        cp.author = "Петро Кравченко"
        cp.title = f"Лабораторна робота №{self.lab}. {theme}"
        cp.subject = "Проєктування інформаційних систем"
        cp.language = "uk-UA"

    # -- sections
    def first_body_section(self):
        """Title page = section 1 (no page number); body starts a new section."""
        # drop trailing empty paragraph of the template so it does not spill to page 2
        last = self.doc.paragraphs[-1]
        if not last.text.strip() and len(self.doc.paragraphs) > 1:
            last._element.getparent().remove(last._element)
        # NB: add_section() turns the template's only (sentinel) sectPr into the *new* section and
        # copies it into the last title paragraph (which becomes section 1 = title page).
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        set_section_geometry(sec, PORTRAIT, WD_ORIENT.PORTRAIT)
        # Body: page number bottom centre; title page is still counted, so the body starts at "2".
        sec.different_first_page_header_footer = False
        sec.footer.is_linked_to_previous = False
        sec.header.is_linked_to_previous = False
        fp = sec.footer.paragraphs[0]
        fmt_par(fp, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, spacing=1.0)
        add_page_number_field(fp)
        # Title page: explicit "Different First Page" with deliberately empty first-page AND default
        # headers/footers — no editor can inherit or show a page number there.
        title = self.doc.sections[0]
        title.different_first_page_header_footer = True
        for hf in (title.first_page_footer, title.footer, title.first_page_header, title.header):
            hf.is_linked_to_previous = False
            for par in hf.paragraphs:
                for r in list(par.runs):
                    r._r.getparent().remove(r._r)

    def switch(self, orientation):
        if orientation == self.orientation:
            return
        sec = self.doc.add_section(WD_SECTION.NEW_PAGE)
        if orientation == "landscape":
            set_section_geometry(sec, LANDSCAPE, WD_ORIENT.LANDSCAPE)
        else:
            set_section_geometry(sec, PORTRAIT, WD_ORIENT.PORTRAIT)
        self.orientation = orientation

    # -- blocks
    @staticmethod
    def number_dot(text):
        """Department rule: a dot after every section number — "5 Назва" / "1.1 Назва" → "5. Назва" / "1.1. Назва"."""
        return re.sub(r"^(\d+(?:\.\d+)*)\.?\s+", r"\1. ", text.strip())

    def _heading(self, text, new_page, align, indent, before, after, bold=True, italic=False, upper=False, level=1):
        self.switch("portrait")
        p = self.doc.add_paragraph()
        p.style = style_by_name(self.doc, f"Heading {level}", outline=level - 1)   # indexed by the TOC field
        fmt_par(p, align=align, indent=indent, before=before, after=after, keep_next=True)
        # page_break_before (not an empty break paragraph) — never produces a blank page after a section break
        p.paragraph_format.page_break_before = bool(new_page or self._break_next)
        self._break_next = False
        txt = self.number_dot(self.resolve(text))
        txt = txt.upper() if upper else txt
        add_text(p, txt, bold=bold, italic=italic)
        if level <= 2:                             # bookmark target for TOC hyperlinks / PAGEREF
            self._bm += 1
            bm = f"_Toc{self._bm:06d}"
            bs = OxmlElement("w:bookmarkStart"); bs.set(qn("w:id"), str(self._bm)); bs.set(qn("w:name"), bm)
            be = OxmlElement("w:bookmarkEnd"); be.set(qn("w:id"), str(self._bm))
            first_run = p._p.find(qn("w:r"))
            first_run.addprevious(bs)
            p._p.append(be)
            self.toc_entries.append((level, txt, bm))

    def h1(self, text, new_page=False):
        self._heading(text, new_page, WD_ALIGN_PARAGRAPH.CENTER, 0, 12, 12, upper=True, level=1)

    def h2(self, text, new_page=False):
        self._heading(text, new_page, WD_ALIGN_PARAGRAPH.LEFT, 1.25, 6, 6, level=2)

    def h3(self, text, new_page=False):
        self._heading(text, new_page, WD_ALIGN_PARAGRAPH.LEFT, 1.25, 6, 0, italic=True, level=3)

    def para(self, text):
        self.switch("portrait")
        for chunk in str(text).strip().split("\n\n"):
            p = self.doc.add_paragraph()
            fmt_par(p)
            add_text(p, self.resolve(" ".join(chunk.split())))

    def lst(self, items, numbered):
        self.switch("portrait")
        for k, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            fmt_par(p)
            marker = f"{k}) " if numbered else "– "
            add_text(p, marker + self.resolve(it))

    def table(self, t):
        """Table with ДСТУ caption. `continue_at: [i, ...]` (0-based data-row indices) splits a long table into
        explicit chunks; each further chunk starts on a new page under "Продовження таблиці N.k" with the header
        repeated — robust in every editor (OnlyOffice ignores cantSplit/keep-with-next inside tables)."""
        self.switch(t.get("orientation", self.orientation if t.get("keep_orientation") else "portrait"))
        self._tab_counter = getattr(self, "_tab_counter", 0) + 1
        num = self._tab_counter
        header, rows = t["header"], t["rows"]
        for i, r in enumerate(rows, 1):
            if len(r) != len(header):
                sys.exit(f"Table '{t['caption']}' row {i} has {len(r)} cells, expected {len(header)}")
        cuts = [0] + sorted(t.get("continue_at", [])) + [len(rows)]
        for ci, (lo, hi) in enumerate(zip(cuts, cuts[1:])):
            if ci == 0:
                caption = f"Таблиця {self.lab}.{num} — {self.resolve(t['caption'])}"
                probe = f"Таблиця {self.lab}.{num} —"
            else:
                caption = probe = f"Продовження таблиці {self.lab}.{num}"
            key = f"{num}.{ci}"
            self._table_chunk(t, header, rows[lo:hi], caption, probe, key,
                              new_page=(ci > 0 or key in self.break_tables or t.get("new_page", False) and ci == 0))

    def _table_chunk(self, t, header, rows, caption, probe, key, new_page):
        cap = self.doc.add_paragraph()
        fmt_par(cap, align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, spacing=1.5, before=6, keep_next=True)
        cap.paragraph_format.page_break_before = bool(new_page)
        add_text(cap, caption)
        size = t.get("font", 12)
        tbl = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders(tbl)
        geo = LANDSCAPE if self.orientation == "landscape" else PORTRAIT
        avail = geo["w"] - geo["left"] - geo["right"]
        widths = t.get("widths") or [1] * len(header)
        widths = [avail * w / sum(widths) for w in widths]
        for j, h in enumerate(header):
            cell_text(tbl.rows[0].cells[j], self.resolve(h), size, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        repeat_table_header(tbl.rows[0])           # <w:tblHeader/> — header repeats on every page
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r):
                val = [self.resolve(x) for x in v] if isinstance(v, list) else self.resolve(v)
                cell_text(tbl.rows[i].cells[j], val, size)
        for row in tbl.rows:                       # <w:cantSplit/> on every row incl. header
            no_split_row(row)
            for j, c in enumerate(row.cells):
                c.width = Cm(widths[j])
        # Keep-together (Word): chain every row to the next for chunks <= 30 rows; the layout loop in main()
        # enforces the same in OnlyOffice by moving a split chunk to a new page.
        keep = t.get("keep_together", len(rows) <= 30)
        last = " ".join(" ".join(x) if isinstance(x, list) else str(x) for x in rows[-1]) if rows else ""
        self.table_meta.append((key, keep, probe, self.resolve(last), " ".join(self.resolve(h) for h in header)))
        chain = tbl.rows[:-1] if keep else tbl.rows[:2]
        for row in chain:
            for c in row.cells:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
        after = self.doc.add_paragraph()
        fmt_par(after, indent=0, spacing=1.0)

    def figure(self, f):
        src = (ROOT / f["src"]).resolve()
        layout = {}
        if src.suffix in (".mmd", ".bpmn", ".puml"):
            png = src.with_suffix(".png")
            lj = src.with_suffix(".layout.json")
            tool = {".mmd": "tools/mmd_render.py", ".bpmn": "tools/bpmn_model.py",
                    ".puml": "tools/puml_render.py"}[src.suffix]
            if not png.exists() or not lj.exists() or png.stat().st_mtime < src.stat().st_mtime:
                r = subprocess.run([sys.executable, str(ROOT / tool), str(src)],
                                   capture_output=True, text=True)
                if r.returncode != 0:
                    sys.exit(f"Diagram {src.name} failed legibility/render:\n{r.stdout}{r.stderr}")
            layout = json.loads(lj.read_text(encoding="utf-8"))
            img = png
        else:
            img = src
            lj = src.with_suffix(".layout.json")        # pre-rendered diagram (e.g. IDEF) with its layout
            if lj.exists():
                layout = json.loads(lj.read_text(encoding="utf-8"))
        orient = f.get("orientation", "auto")
        if orient == "auto":
            orient = layout.get("orientation", "portrait")
        self.switch(orient)
        width = f.get("width_cm") or (layout[orient]["width_cm"] if orient in layout else
                                      (16.5 if orient == "portrait" else 25.7))
        p = self.doc.add_paragraph()
        fmt_par(p, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, spacing=1.0, before=6, keep_next=True)
        p.add_run().add_picture(str(img), width=Cm(width))
        self._fig_counter = getattr(self, "_fig_counter", 0) + 1
        cap = self.doc.add_paragraph()
        fmt_par(cap, align=WD_ALIGN_PARAGRAPH.CENTER, indent=0, spacing=1.5, before=6, after=6)
        add_text(cap, f"Рис. {self.lab}.{self._fig_counter} — {self.resolve(f['caption'])}")
        if f.get("note"):
            n = self.doc.add_paragraph()
            fmt_par(n, align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, spacing=1.0, after=6)
            add_text(n, self.resolve(f["note"]), size=12, italic=True)

    def pagebreak(self):
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)

    # -- main
    def build(self):
        self.fill_title()
        self.first_body_section()
        if self.spec.get("toc", True):
            self.toc()
        for b in self.spec["blocks"]:
            key = next(iter(b)); val = b[key]
            np_ = bool(b.get("new_page"))
            if key == "h1": self.h1(val, np_)
            elif key == "h2": self.h2(val, np_)
            elif key == "h3": self.h3(val, np_)
            elif key == "p": self.para(val)
            elif key == "bullets": self.lst(val, False)
            elif key == "numbered": self.lst(val, True)
            elif key == "table": self.table(val)
            elif key == "figure": self.figure(val)
            elif key == "pagebreak": self.pagebreak()
            else: sys.exit(f"Unknown block type: {key}")
        out = (ROOT / self.spec["output"]).resolve()
        out.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(out))
        return out


WORD = re.compile(r"[A-Za-zА-ЯІЇЄҐа-яіїєґ’']{5,}")


def split_tables(pdf, meta):
    """Return keys of keep-together table chunks whose last row is not on the caption page (PDF as rendered
    by OnlyOffice, the user's editor — it ignores keep-with-next/cantSplit inside tables)."""
    n = int(subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
            .split("Pages:")[1].split()[0])
    pages = [subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), str(pdf), "-"],
                            capture_output=True, text=True).stdout for i in range(1, n + 1)]
    # -layout keeps table rows on one line (plain mode reads tables column by column)
    pages_lay = [subprocess.run(["pdftotext", "-layout", "-f", str(i), "-l", str(i), str(pdf), "-"],
                                capture_output=True, text=True).stdout for i in range(1, n + 1)]
    bad, used = [], {}
    for key, keep, probe, last, head in meta:
        start = used.get(probe, 0)
        pg = next((i for i in range(start, n) if probe in pages[i]), None)
        if pg is None:
            continue
        used[probe] = pg + 1
        if not keep:
            continue
        split = False
        words = set(WORD.findall(last))
        if words:   # (a) last row's words missing from the caption page
            on_page = set(WORD.findall(pages[pg].split(probe, 1)[1]))
            split = len(words & on_page) / len(words) < 0.8
        if not split and head and pg + 1 < n:  # (b) header row repeated at the top of the next page
            nxt = pages_lay[pg + 1].lstrip()
            if not nxt.startswith(("Таблиця", "Продовження")):
                top = nxt[:len(head) * 3 + 200]
                compact = lambda t: re.sub(r"[\s\-]+", "", t)   # header cell wrapped mid-word
                toks = lambda t: set(re.findall(r"[A-Za-zА-ЯІЇЄҐа-яіїєґ’'\-]{2,}", t))
                ht = toks(head)                                     # header cells wrapped onto 2+ lines
                split = compact(head)[:40] in compact(top) or (
                    len(ht) >= 2 and len(ht & toks(top)) / len(ht) >= 0.8)
        if split:
            bad.append(key)
    return bad


def heading_pages(pdf, entries):
    """Find the page of every TOC heading in the rendered PDF (searching in document order after ЗМІСТ)."""
    n = int(subprocess.run(["pdfinfo", str(pdf)], capture_output=True, text=True).stdout
            .split("Pages:")[1].split()[0])
    norm = lambda t: " ".join(t.split())
    pages = [norm(subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), str(pdf), "-"],
                                 capture_output=True, text=True).stdout) for i in range(1, n + 1)]
    cur = next((i for i, t in enumerate(pages) if "ЗМІСТ" in t), 0) + 1
    found = {}
    for _, txt, bm in entries:
        key = norm(txt)[:60]
        pg = next((i for i in range(cur, n) if key in pages[i]), None)
        if pg is not None:
            found[bm] = pg + 1                     # displayed number == physical page (title page counted)
            cur = pg
    return found


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("spec")
    ap.add_argument("--pdf", action="store_true",
                    help="render PDF + previews and run the table layout loop (recommended)")
    a = ap.parse_args()
    spec_path = Path(a.spec).resolve()
    spec = yaml.safe_load(spec_path.read_text(encoding="utf-8"))
    breaks, pages = set(), {}
    for attempt in range(8):
        rb = ReportBuilder(spec, spec_path.parent, breaks, pages)
        out = rb.build()
        if not a.pdf:
            break
        pdf = out.with_suffix(".pdf")
        subprocess.run([str(ROOT / "tools/docx2pdf.sh"), str(out), str(pdf)], check=True, capture_output=True)
        bad = [b for b in split_tables(pdf, rb.table_meta) if b not in breaks]
        if bad:
            print(f"layout pass {attempt + 1}: tables {bad} split across pages -> start them on a new page")
            breaks |= set(bad)
            continue
        new_pages = heading_pages(pdf, rb.toc_entries)
        if new_pages == pages:
            break
        print(f"layout pass {attempt + 1}: updating ЗМІСТ page numbers")
        pages = new_pages
    print(f"DOCX: {out}")
    if a.pdf:
        rest = split_tables(pdf, rb.table_meta)
        if rest:
            print(f"WARNING: table chunks {rest} still split (taller than a page) — add continue_at to the spec")
        prev = out.parent / "preview"
        prev.mkdir(exist_ok=True)
        for old in prev.glob("page-*.png"):
            old.unlink()
        subprocess.run(["pdftoppm", "-png", "-r", "70", str(pdf), str(prev / "page")], check=True)
        print(f"PDF: {pdf}\nPreviews: {prev}/page-*.png")


if __name__ == "__main__":
    main()
