"""Спільне для практичних робіт «Системи банкової безпеки» (варіант 08).

Шаблони викладача містять аркуші всіх варіантів (префікс «В01 »…«В10 »). Для здачі лишаємо тільки аркуші
свого варіанта: формули посилаються на них за назвами, тому аркуші не перейменовуємо.
"""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "Системи банкової безпеки"
VAR = "08"
STUDENT = "Кравченко"
PREFIX = f"В{VAR} "


def variant_book(n):
    """Завантажити шаблон ПР n і прибрати аркуші інших варіантів та «Зміст»."""
    wb = openpyxl.load_workbook(SRC / f"Практична {n}" / f"ПР_{n}_Шаблони_Усі_варіанти.xlsx")
    for name in list(wb.sheetnames):
        if not name.startswith(PREFIX):
            del wb[name]
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # openpyxl не зберігає обчислені значення — Excel перерахує
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws is wb.active
    return wb


def sheet(wb, name):
    return wb[PREFIX + name]


def put_rows(ws, first_row, rows, cols):
    """rows — список кортежів; cols — літери стовпців, у які пишемо значення (None — пропустити)."""
    for i, row in enumerate(rows):
        for col, val in zip(cols, row):
            if col and val is not None:
                ws[f"{col}{first_row + i}"] = val


def out_path(n, kind="xlsx", what="Варіант"):
    d = SRC / f"Практична {n}" / "Готово"
    d.mkdir(exist_ok=True)
    return d / f"ПР_{n}_{what}_{VAR}_{STUDENT}.{kind}"


def evaluate(path):
    """Перерахувати формули книги (бібліотека formulas) → {('АРКУШ', 'A1'): значення}."""
    import formulas
    xl = formulas.ExcelModel().loads(str(path)).finish()
    sol = xl.calculate()
    out = {}
    for k, v in sol.items():
        # ключ виду "'[file.xlsx]В08 РИЗИКИ'!R5"
        k = str(k)
        if "]" not in k or "!" not in k:
            continue
        sh, cell = k.split("]", 1)[1].rsplit("!", 1)
        val = v.value[0, 0] if hasattr(v, "value") else v
        out[(sh.strip("'").upper(), cell.upper())] = val
    return out


def val(res, sheet_name, cell):
    v = res.get(((PREFIX + sheet_name).upper(), cell.upper()))
    if v is None:
        return None
    s = str(v)
    return "" if s in ("", "empty") else v


def fill_blank(n, answers, strong, fix, date="24.10.2026"):
    """Заповнити копію бланка самоперевірки ПР n: answers — 7 відповідей у порядку питань."""
    import docx
    d = docx.Document(SRC / f"Практична {n}" / f"Практична_{n}_Бланк_перевірки.docx")
    head = {2: f"Перевіряє (прізвище): {STUDENT} (самоперевірка)   Робота студента (прізвище, варіант №): "
               f"{STUDENT}, варіант {VAR}   Дата: {date}",
            5: f"Сильна сторона роботи: {strong}", 6: f"Що виправити насамперед: {fix}"}
    for i, text in head.items():
        p = d.paragraphs[i]
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    t = d.tables[0]
    assert len(t.rows) == 8 and len(answers) == 7
    for row, a in zip(t.rows[1:], answers):
        cell = row.cells[1]
        cell.paragraphs[0].text = a
    out = out_path(n, "docx", "Бланк_перевірки")
    d.save(out)
    return out
