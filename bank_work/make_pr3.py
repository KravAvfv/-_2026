"""ПР3 «Проєкт мережі відділення: зони, VLAN, адреси», варіант 08 → Практична 3/Готово/ (xlsx + pptx + бланк).

В08: блок 10.209.0.0/16, шлюз — остання корисна адреса підмережі; особливі пристрої — мережевий сейф і ноутбуки
підрядників. У шаблоні викладача колонка «Перевірка» аркуша «Зони й адреси» (L5:L13) порожня в усіх варіантах,
хоча «Підсумок» на неї спирається, — відновлюємо її формулою за правилами аркуша «Інструкція».
"""
import math
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Pt

from common import STUDENT, VAR, evaluate, fill_blank, out_path, put_rows, sheet, val, variant_book

A = 209
GROWTH = {"KASA": 15, "MANAGERS": 20, "ABS": 30, "ATM_POS": 20, "IOT": 50, "GUEST": 25, "MGMT": 20}
DEVICE_ZONE = ["KASA", "KASA", "MANAGERS", "ABS", "ATM_POS", "ATM_POS", "IOT", "IOT", "IOT", "GUEST", "MGMT",
               "MGMT", "MGMT", "MGMT", "IOT", "GUEST"]            # рядки 5–20 аркуша «Пристрої»
DEVICES = [4, 2, 6, 3, 2, 8, 12, 5, 5, 80, 2, 3, 2, 1, 1, 2]
VLAN = {"KASA": 10, "MANAGERS": 20, "ABS": 30, "ATM_POS": 40, "IOT": 50, "GUEST": 90, "MGMT": 100}
VNAME = {"KASA": "V10_KASY", "MANAGERS": "V20_MENEDZHERY", "ABS": "V30_ABS", "ATM_POS": "V40_ATM_POS",
         "IOT": "V50_IOT", "GUEST": "V90_GUEST", "MGMT": "V100_MGMT"}
NATIVE = 250
ZONES = ["KASA", "MANAGERS", "ABS", "ATM_POS", "IOT", "GUEST", "MGMT"]  # рядки 5–11 «Зон й адрес»


def plan():
    """Розрахунок за правилами «Інструкції»: потрібно адрес, мінімальний префікс, мережа, шлюз (остання)."""
    out = {}
    for z in ZONES:
        n_dev = sum(c for c, zz in zip(DEVICES, DEVICE_ZONE) if zz == z)
        need = math.ceil(round(n_dev * (1 + GROWTH[z] / 100), 6)) + 1
        prefix = 32 - max(2, math.ceil(math.log2(need + 2) - 1e-6))
        size = 2 ** (32 - prefix)
        out[z] = dict(dev=n_dev, need=need, prefix=prefix, usable=size - 2, net=f"10.{A}.{VLAN[z]}.0",
                      gw=f"10.{A}.{VLAN[z]}.{size - 2}")
    return out


CHECK = ('=IF(C{r}=0,"",IF(E{r}<>N{r},"перерахуйте потрібну кількість адрес",'
         'IF(NOT(AND(ISNUMBER(F{r}),F{r}>=2,F{r}<=254,F{r}=INT(F{r}))),"VLAN — ціле число від 2 до 254",'
         'IF(COUNTIF($F$5:$F$13,F{r})>1,"VLAN повторюється",'
         'IF(SUMPRODUCT(($F$5:$F$13<>"")*(ABS($F$5:$F$13-F{r})<5))>1,"між VLAN сусідніх зон проміжок менше 5",'
         'IF(AND($C$17<>"",$C$18<>"",F{r}>=$C$17,F{r}<=$C$18),"VLAN у зарезервованому діапазоні",'
         'IF(NOT(AND(ISNUMBER(I{r}),I{r}>=24,I{r}<=30)),"префікс — число від 24 до 30",'
         'IF(I{r}>O{r},"підмережа замала для потрібної кількості адрес",'
         "IF(OR(P{r}<>10,Q{r}<>'В08 Картка'!$B$5,R{r}<>F{r}),\"мережа має бути 10.A.<VLAN>.x\","
         'IF(MOD(S{r},2^(32-I{r}))<>0,"адреса мережі не на межі підмережі",'
         "IF(OR(U{r}<>1,T{r}<>IF(ISNUMBER(SEARCH(\"остання\",'В08 Картка'!$B$6)),S{r}+2^(32-I{r})-2,S{r}+1)),"
         '"шлюз не за правилом варіанта","ОК")))))))))))')

DRAFT_LOG = [
    ("3", "VLAN 1",
     "Термінали АБС у VLAN 1 — це VLAN за замовчуванням, правила забороняють його використовувати. Треба окремий "
     "номер із проміжком (у мене 30) і мережа 10.209.30.0/29, шлюз .6."),
    ("5", "Адреса мережі не на межі / шлюз не за правилом",
     "/26 має починатися з адреси, кратної 64: .32 — не межа. Правильно 10.209.170.0/26; шлюз — остання корисна "
     "адреса .62, а не перша .33."),
    ("7", "Мережа не відповідає VLAN",
     "Для VLAN 190 мережа має бути 10.209.190.0/28, а колега скопіював 10.209.140.0 з рядка «Менеджери» — "
     "дві зони з однією підмережею. Шлюз тоді 10.209.190.14."),
]

JUST_DEV = [
    ("IOT", "Сейф — пристрій без користувача, як замки й датчики, тож кладу його в IoT. У зоні касирів він був би "
     "доступний з будь-якого ПК каси, а в гостьовій — узагалі з Інтернету; в IoT доступ до нього дає лише ACL "
     "із MGMT."),
    ("GUEST", "Ноутбуки підрядників — чужі пристрої, яким банк не довіряє, тому вони в гостьовій зоні. Якби "
     "поклали в MGMT, чужий ноутбук мав би прямий доступ до всіх комутаторів; налаштування — лише через "
     "станцію адміністратора під наглядом."),
]
VLAN_LOGIC = ("VLAN кратні 10: 10–40 — робочі зони з грошима й даними, 50 — пристрої без користувача, 90 — "
              "гості (в кінці, найменш довірена), 100 — керування. Третій октет = VLAN, тож з адреси одразу видно "
              "зону; між номерами 10 вільних — нову зону (напр., 60) можна додати без перенумерації. Native VLAN "
              f"{NATIVE} не використовує жодна зона.")
INJECT = ("КРОК 4 НЕ ВИКОНАНО: файл «ПР_3_Нова_обставина_Варіант_08» видає викладач після кроків 1–3; рядки NEW1/"
          "NEW2, обмеження ГО і слайд 2 заповнюються після його отримання.")
SELF = ("Порівняв свій розрахунок адрес із чернеткою колеги — кількості збіглися (IoT 36, гості 104), отже "
        "особливі пристрої розкладено так само. Перевірив вирівнювання кожної мережі за префіксом і шлюз .14/.6/"
        ".62/.126 (остання корисна).")


def build_xlsx():
    wb = variant_book(3)
    sheet(wb, "Картка")["B3"] = STUDENT
    ws = sheet(wb, "Пристрої")
    for i, z in enumerate(DEVICE_ZONE):
        ws[f"E{5 + i}"] = z
    put_rows(sheet(wb, "Чернетка колеги"), 15, DRAFT_LOG, "ABC")
    za = sheet(wb, "Зони й адреси")
    p = plan()
    for i, z in enumerate(ZONES):
        r = 5 + i
        assert za[f"A{r}"].value == z
        d = p[z]
        for col, v in zip("EFGHIJ", (d["need"], VLAN[z], VNAME[z], d["net"], d["prefix"], d["gw"])):
            za[f"{col}{r}"] = v
    for r in range(5, 14):
        za[f"L{r}"] = CHECK.format(r=r)
    za["C20"] = NATIVE
    ob = sheet(wb, "Обґрунтування")
    put_rows(ob, 5, JUST_DEV, "BC")
    ob["A16"], ob["A19"], ob["A22"] = VLAN_LOGIC, INJECT, SELF
    out = out_path(3)
    wb.save(out)
    return out


# --------------------------------------------------------------------------- PowerPoint: схема зон
NAVY = RGBColor(0x1F, 0x38, 0x64)
ZCOL = {"KASA": "FCE4D6", "MANAGERS": "E2EFDA", "ABS": "FCE4D6", "ATM_POS": "FCE4D6", "IOT": "EDEDED",
        "GUEST": "FFF2CC", "MGMT": "DDEBF7"}
ZTXT = {"KASA": "Каси\n(4 РМ, 2 принтери)", "MANAGERS": "Менеджери\n(6 РМ)", "ABS": "Термінали АБС\n(3)",
        "ATM_POS": "Банкомати й POS\n(2 + 8)", "IOT": "IoT: камери, замки,\nдатчики, сейф",
        "GUEST": "Гостьова Wi-Fi\n(80 + 2 ноутб. підрядн.)", "MGMT": "Керування:\nкомутатори, AP, МЕ,\nстанція адм."}


def _box(sl, x, y, w, h, text, fill="E8EEF7", size=11, bold=False):
    s = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Cm(x), Cm(y), Cm(w), Cm(h))
    s.fill.solid(); s.fill.fore_color.rgb = RGBColor.from_string(fill)
    s.line.color.rgb = NAVY; s.line.width = Pt(1.25)
    tf = s.text_frame; tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Cm(0.08))
    for i, line in enumerate(text.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = PP_ALIGN.CENTER
        r = para.add_run(); r.text = line
        r.font.size = Pt(size); r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.bold = bold and i == 0
    return s


def _line(sl, x1, y1, x2, y2, w=1.75):
    c = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    c.line.color.rgb = NAVY; c.line.width = Pt(w)


def _text(sl, x, y, w, h, text, size=12, bold=False, color=NAVY, align=PP_ALIGN.CENTER):
    t = sl.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = t.text_frame; tf.word_wrap = True
    tf.text = text
    for para in tf.paragraphs:
        para.alignment = align
        for r in para.runs:
            r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color


def build_pptx():
    p = plan()
    prs = Presentation()
    prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
    sl = prs.slides.add_slide(prs.slide_layouts[6])
    _text(sl, 0.5, 0.3, 32.8, 1.2, f"ПР 3. Зони, VLAN і адреси відділення — варіант {VAR} ({STUDENT}), "
          f"блок 10.{A}.0.0/16", 20, True)
    _box(sl, 13.4, 1.9, 7.0, 1.5, "Головний офіс / Інтернет", "FFFFFF", 12, True)
    _box(sl, 12.4, 4.3, 9.0, 1.9, "Маршрутизатор і міжмережевий екран\nшлюзи зон = остання корисна адреса",
         "DDEBF7", 12, True)
    _line(sl, 16.9, 3.4, 16.9, 4.3)
    _box(sl, 12.4, 7.2, 9.0, 1.7, f"Комутатор відділення\ntrunk до МЕ, native VLAN {NATIVE}", "DDEBF7", 12, True)
    _line(sl, 16.9, 6.2, 16.9, 7.2, 3)
    _text(sl, 17.1, 6.25, 6.0, 0.9, "802.1Q trunk", 10, False, RGBColor(0x59, 0x59, 0x59), PP_ALIGN.LEFT)
    n, w, gap = len(ZONES), 4.35, 0.32
    x0 = (33.867 - (n * w + (n - 1) * gap)) / 2
    ybus = 10.0
    _line(sl, x0 + w / 2, ybus, x0 + (n - 1) * (w + gap) + w / 2, ybus)
    _line(sl, 16.9, 8.9, 16.9, ybus)
    for i, z in enumerate(ZONES):
        x = x0 + i * (w + gap)
        _line(sl, x + w / 2, ybus, x + w / 2, 10.8)
        d = p[z]
        _box(sl, x, 10.8, w, 5.6, f"VLAN {VLAN[z]} · {z}\n{ZTXT[z]}\n\n{d['net']}/{d['prefix']}\n"
             f"шлюз .{d['gw'].rsplit('.', 1)[1]}\n{d['usable']} адрес (треба {d['need']})", ZCOL[z], 11, True)
    _text(sl, 0.6, 16.8, 32.6, 1.8, "Між зонами трафік іде лише через міжмережевий екран (ACL): гості — тільки в "
          "Інтернет; IoT і сейф — доступні лише з MGMT; каси, АБС і POS — до шлюзу ГО. Ноутбуки підрядників — у "
          "гостьовій зоні, налаштування обладнання — зі станції адміністратора.", 11, False,
          RGBColor(0, 0, 0), PP_ALIGN.LEFT)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    _text(s2, 0.5, 0.3, 32.8, 1.2, f"Схема після нової обставини — варіант {VAR}", 20, True)
    _box(s2, 6.0, 7.0, 21.8, 4.0, "Слайд заповнюється після отримання файлу «ПР_3_Нова_обставина_Варіант_08» "
         "(крок 4): нові зони NEW1/NEW2 або зміни VLAN.", "FFF2CC", 16)
    out = out_path(3, "pptx").with_name(f"ПР_3_Схема_Варіант_{VAR}_{STUDENT}.pptx")
    prs.save(out)
    return out


def verify(path):
    res = evaluate(path)
    ok = True
    for i, z in enumerate(ZONES):
        r = 5 + i
        got = [val(res, "Зони й адреси", f"{c}{r}") for c in "CEHIJKLNO"]
        print(f"  {z:<9} пристроїв={got[0]} треба={got[1]} {got[2]}/{got[3]} шлюз {got[4]} корисних={got[5]} "
              f"N={got[7]} O={got[8]} → {got[6]}")
        if got[6] != "ОК":
            ok = False
    print("  native:", val(res, "Зони й адреси", "L20"))
    for r in range(4, 12):
        print("  Підсумок", val(res, "Підсумок", f"A{r}"), "=", val(res, "Підсумок", f"B{r}"))
    if val(res, "Підсумок", "B11") != "Можна здавати":
        ok = False
    for r in range(5, 21):
        f = val(res, "Пристрої", f"F{r}")
        if f not in ("ОК", "рішення поясніть на аркуші «Обґрунтування»"):
            print("  Пристрої", r, f); ok = False
    print("PASS" if ok else "FAIL")
    return ok


def selftest_check(path):
    """Відновлена перевірка має ловити помилки чернетки колеги (VLAN 1, не та межа, чужа мережа)."""
    import openpyxl
    import tempfile
    wb = openpyxl.load_workbook(path)
    za = wb[f"В{VAR} Зони й адреси"]
    za["F7"], za["H7"], za["J7"] = 1, f"10.{A}.1.0", f"10.{A}.1.6"          # ABS у VLAN 1
    za["H9"], za["J9"] = f"10.{A}.50.32", f"10.{A}.50.33"                   # IOT: не межа /26
    za["H11"] = f"10.{A}.20.0"                                              # MGMT: чужа мережа
    za["I5"] = 29                                                           # KASA: /29 замала
    with tempfile.TemporaryDirectory() as d:
        t = f"{d}/t.xlsx"; wb.save(t)
        res = evaluate(t)
    got = {r: val(res, "Зони й адреси", f"L{r}") for r in (5, 7, 9, 11)}
    print("  самоперевірка формули:", got)
    exp = {5: "підмережа замала", 7: "VLAN — ціле", 9: "адреса мережі не на межі", 11: "мережа має бути"}
    return all(str(got[r]).startswith(exp[r]) for r in exp)


BLANK = [
    "Усі 16 груп мають зону; не знайшов пристрою не на своєму місці.",
    "Так: сейф — IoT (у касах був би доступний з ПК касирів), ноутбуки підрядників — гості (у MGMT мали б прямий "
    "доступ до комутаторів).",
    "IoT: (12 + 5 + 5 + 1) × 1,5 = 34,5 → 35, + шлюз = 36 — збігається з колонкою «Потрібно адрес».",
    "Не знайшов: у кожної зони мінімальний префікс, що вміщує потрібні адреси (напр., гості 104 → /25 на 126).",
    "VLAN 10, 20, 30, 40, 50, 90, 100: проміжки ≥ 10, VLAN 1 немає, логіка — від зон з грошима до гостей і "
    "керування.",
    f"Усі мережі на межі префікса, шлюз — остання корисна адреса; native VLAN {NATIVE} окремий.",
    "Ще не виконано: файл «Нова обставина» видає викладач після кроків 1–3.",
]

if __name__ == "__main__":
    x = build_xlsx(); print(x)
    print(build_pptx())
    print(fill_blank(3, BLANK, "Розрахунок адрес збігся з чернеткою колеги, кожна мережа вирівняна, VLAN "
                     "пронумеровані із запасом.", "Опрацювати «Нову обставину» (рядки NEW1/NEW2)."))
    ok = verify(x)
    ok = selftest_check(x) and ok
    print("ALL PASS" if ok else "ALL FAIL")
    sys.exit(0 if ok else 1)
