#!/usr/bin/env python3
"""Тези доповіді на Міжнародну науково-технічну Internet-конференцію НУХТ «Сучасні методи, інформаційне, програмне
та технічне забезпечення систем керування організаційно-технічними та технологічними комплексами».

Завдання (фото «НУХТ/photo_2026-10-24_12-56-37.jpg»): УДК 004.9, назва «Архітектура …», автори, КНУ; 1) абзац
загального вступу; 2) абзаци аналізу області застосування (вимоги з висновків ПР1 і ПР2); 3) рисунок архітектури;
4) опис архітектури; 5) висновки про майбутнє проєктне рішення і як саме реалізовуватиметься система;
6) джерела з посиланнями в абзацах 1–2. Оформлення — точно за шаблоном «НУХТ/Example_2025.doc».

Шаблон: стилі беруться з самого Example_2025.doc (конвертованого в .docx через x2t), тіло замінюється текстом тез.
Run:  .venv/bin/python tspp_work/nuft/make_theses.py
"""
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(ROOT / "tools"))
sys.path.insert(0, str(ROOT / "tspp_work" / "pract2"))
from idef_model import render  # noqa: E402
from model_p2 import G1, G2, G3  # noqa: E402

SRC = ROOT / "Технології створення програмних продуктів" / "НУХТ"
OUT = SRC / "Готово"
BASE = "Кравченко_П"          # оргкомітет: «файл ПІБ.doc з текстом тези, у який вставлено рисунки»
X2T = Path("/opt/onlyoffice/desktopeditors/converter")
FONTS = Path.home() / ".local/share/onlyoffice/desktopeditors/data/fonts/AllFonts.js"
FIG_W = 16.5                  # см; поле набору 17 см (A4, поля по 20 мм)

UDC = "УДК 004.9"
TITLE = ("Архітектура розподіленої IoT-системи цілевказівки та координації дронів для перехоплення "
         "повітряних цілей")
AUTHORS = "П.О. Кравченко, О.В. Кравченко"
ORG = "Київський національний університет імені Тараса Шевченка"

# Компактний варіант рисунка архітектури з ПР2 (tspp_work/pract2/model_p2.py: ARCH) — ті самі рівні та
# компоненти, коротші підписи, щоб текст на ширині 16,5 см був ≥ 10 пт.
FIG = {
    "kind": "columns", "colw": 162, "bw": 136, "rh": 56, "gap": 14,
    "columns": [
        {"title": "Пристрої", "layer": "Рівень датчиків", "fill": G1, "items": [
            {"id": "AC", "name": "Акустичні вузли", "row": 0},
            {"id": "RD", "name": "Малогабаритні РЛС", "row": 1},
            {"id": "EO", "name": "Оптико-електронні пости", "row": 2},
            {"id": "GCS", "name": "Перехоплювачі та НСУ", "row": 3},
        ]},
        {"title": "Зв’язок", "layer": "Мережевий рівень", "fill": G2, "items": [
            {"id": "GW", "name": "Периферійний шлюз: MQTT, злиття треків, буфер", "row": 0, "span": 3},
            {"id": "CH", "name": "Канали зв’язку, VPN", "row": 3},
        ]},
        {"title": "IoT-ядро", "layer": "Рівень обробки", "fill": G3, "items": [
            {"id": "BUS", "name": "Шина подій", "row": 0},
            {"id": "SRV", "name": "Мікросервіси", "row": 1},
            {"id": "DB", "name": "PostgreSQL, TimescaleDB", "row": 2},
            {"id": "API", "name": "API-шлюз: REST, WebSocket", "row": 3},
        ]},
        {"title": "Застосунки", "layer": "Рівень застосунків", "fill": G1, "items": [
            {"id": "CPW", "name": "АРМ командного пункту", "row": 0.5},
            {"id": "TAB", "name": "Планшет розрахунку", "row": 1.5},
            {"id": "EXT", "name": "Зовнішні системи", "row": 2.5},
        ]},
    ],
    "links": [("AC", "GW"), ("RD", "GW"), ("EO", "GW"), ("GCS", "CH", {"both": True}),
              ("GW", "CH", {"both": True}), ("CH", "BUS", {"both": True}),
              ("BUS", "SRV", {"both": True}), ("SRV", "DB", {"both": True}), ("DB", "API", {"both": True}),
              ("API", "CPW", {"both": True}), ("API", "TAB", {"both": True}), ("API", "EXT", {"both": True})],
}

# ---------------------------------------------------------------- текст тез
# [n] — посилання на REFS (нумерація за першим згадуванням; усі посилання — в абзацах вступу та аналізу)
INTRO = [
    "Масове застосування малих безпілотних літальних апаратів перетворило захист повітряного простору на "
    "задачу обробки великих потоків різнотипних даних у реальному часі. Протидію забезпечує розподілена "
    "система, що зводить дані багатьох датчиків в єдину картину обстановки й координує дії людей. Для таких "
    "систем доцільно застосувати підхід інтернету речей (Internet of Things, IoT) з чотирирівневою "
    "архітектурою: рівні датчиків, мережевий, обробки даних і застосунків [1]. Цю модель використано в "
    "системах медичного моніторингу [1], керування ветеринарною клінікою [2] та насосною станцією "
    "водопостачання [3].",
]
ANALYSIS = [
    "Існуючі рішення закривають лише частину задачі. Система ситуаційної обізнаності DELTA побудована модульно "
    "навколо єдиної цифрової карти й має мобільний клієнт з роботою без зв’язку [4]. Платформа "
    "DedroneTracker.AI зводить дані радіочастотних, радіолокаційних, відео- та акустичних датчиків у "
    "спільному рушії злиття [5]. ThingsBoard Edge забезпечує автономну роботу периферійного вузла з "
    "синхронізацією після відновлення зв’язку [6]. Жодне з рішень не поєднує мережу недорогих датчиків, "
    "периферійні обчислення та координацію завдань перехоплення на власній інфраструктурі.",
    "Для системи сформульовано 22 функціональні та 16 нефункціональних вимог. Визначальними для архітектури є "
    "затримка «датчик — екран оператора» до 2 с, доставка завдання на планшет розрахунку до 1 с, автономна "
    "робота позиції без зв’язку з центром щонайменше 4 год, взаємна автентифікація пристроїв, повний журнал "
    "аудиту та принцип «людина в контурі»: жодне застосування засобу перехоплення не відбувається без санкції "
    "чергового.",
]
ARCH_TEXT = [
    "Архітектуру IoT-платформи «Щит-Лінк» побудовано за чотирирівневою моделлю (Рис. 1). Рівень датчиків "
    "утворюють акустичні вузли, малогабаритні РЛС, оптико-електронні пости та наземні станції управління (НСУ) "
    "перехоплювачами, що підключаються стандартними протоколами.",
]
ARCH_TEXT2 = [
    "На мережевому рівні периферійний шлюз позиції приймає дані через брокер MQTT, перетворює протоколи "
    "пристроїв у єдиний формат, локально зводить виявлення в треки й буферизує дані на час втрати зв’язку. "
    "Обмін із ядром іде резервованими каналами (mesh-радіо, LTE, супутник) у VPN з автентифікацією mTLS.",
    "Рівень обробки даних — IoT-ядро у власних центрах обробки даних у гібридному стилі: подієво-орієнтована "
    "взаємодія через шину подій і дев’ять мікросервісів (треки, місії, ресурси, аудит тощо). Дані "
    "зберігаються в PostgreSQL з PostGIS і TimescaleDB для телеметрії. Рівень застосунків через API-шлюз "
    "(REST, WebSocket) обслуговує веб-АРМ командного пункту, планшет розрахунку та зовнішні системи.",
]
CONCLUSION = [
    "Отже, архітектура поєднує переваги аналогів: периферійні обчислення дають автономність позицій, "
    "шина подій і мікросервіси — затримку до 2 с і масштабування, а підтвердження людиною — контроль над "
    "кожним застосуванням. Систему буде реалізовано інкрементно "
    "(V-модель у кожному інкременті, випуск за практиками DevOps): спочатку ядро обліку на PostgreSQL з "
    "журналом аудиту й веб-інтерфейсом, далі шлюз позиції, сервіси реального часу "
    "й мобільний застосунок, наприкінці — зовнішні інтеграції.",
]

ACC = "[Дата звернення 4 Жовтень 2026]"
REFS = [
    "Zinko, I., Kravchenko, O. та Syvoglaz, D., 2024. Designing an Internet of Things solution for monitoring "
    "vital signs. Technology Audit and Production Reserves, 2(2(76)). "
    "doi:10.15587/2706-5448.2024.300745.",
    "Васюта, А. та Кравченко, О., 2025. IoT система керування ветеринарною клінікою. В: Інформаційні системи "
    "та технології: результати і перспективи (IST 2025): матеріали 2-ї Міжнар. наук.-практ. конф. Київ: "
    "ФІТ КНУ імені Тараса Шевченка, с.36-38.",
    "Легойда, Ю.В. та Кравченко, О.В., 2024. Автоматизована IoT система управління насосної станції "
    "водопостачання селища Теплик. Таврійський науковий вісник. Серія: Технічні науки, 2, с.73-87.",
    "Міністерство оборони України, 2026. DELTA: Open Wiki. [online] Доступно: "
    f"<https://delta.mil.gov.ua/open-wiki/en/> {ACC}.",
    "Dedrone, 2026. DedroneTracker.AI: drone detection software. [online] Доступно: "
    f"<https://www.dedrone.com/products/drone-detection-software> {ACC}.",
    "ThingsBoard, 2026. ThingsBoard Edge. [online] Доступно: "
    f"<https://thingsboard.io/products/thingsboard-edge/> {ACC}.",
]
CAPTION = "Рис. 1. Архітектура IoT-платформи «Щит-Лінк»"


# ---------------------------------------------------------------- збирання
def x2t(src, dst, fmt):
    """Конвертація OnlyOffice x2t (65 = docx, 66 = doc, 513 = pdf); ASCII-копії через кирилицю в шляхах."""
    with tempfile.TemporaryDirectory() as t:
        t = Path(t)
        a, b = t / f"in{src.suffix}", t / f"out{dst.suffix}"
        shutil.copy(src, a)
        (t / "task.xml").write_text(
            f'<?xml version="1.0" encoding="utf-8"?><TaskQueueDataConvert><m_sFileFrom>{a}</m_sFileFrom>'
            f'<m_sFileTo>{b}</m_sFileTo><m_nFormatTo>{fmt}</m_nFormatTo><m_sFontDir>/usr/share/fonts</m_sFontDir>'
            f'<m_sAllFontsPath>{FONTS}</m_sAllFontsPath><m_bIsNoBase64>true</m_bIsNoBase64></TaskQueueDataConvert>',
            encoding="utf-8")
        subprocess.run(["./x2t", str(t / "task.xml")], cwd=X2T, env={"LD_LIBRARY_PATH": "."}, check=True,
                       capture_output=True)
        if not b.exists() or not b.stat().st_size:
            raise SystemExit(f"x2t: no output for {dst.name}")
        shutil.copy(b, dst)


def template(tmp):
    """Example_2025.doc → .docx; jc="end" (ISO-варіант) → "right", щоб python-docx і Word 2003 його читали."""
    docx = tmp / "example.docx"
    x2t(SRC / "Example_2025.doc", docx, 65)
    fixed = tmp / "base.docx"
    import zipfile
    with zipfile.ZipFile(docx) as zi, zipfile.ZipFile(fixed, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it)
            if it.filename in ("word/document.xml", "word/styles.xml"):
                data = data.replace(b'w:val="end"', b'w:val="right"').replace(b'w:val="start"', b'w:val="left"')
            zo.writestr(it, data)
    return fixed


def style(doc, sid):
    return next(s for s in doc.styles if s.style_id == sid)


def build(fig_png):
    with tempfile.TemporaryDirectory() as t:
        doc = Document(template(Path(t)))
    body = doc.element.body
    for el in list(body):
        if not el.tag.endswith("sectPr"):
            body.remove(el)
    S = {k: style(doc, v) for k, v in dict(h1="Heading1", h2="Heading2", auth="UserStyle_0", org="UserStyle_1",
                                            text="UserStyle_3", cap="UserStyle_4", ref="UserStyle_10").items()}

    p = doc.add_paragraph(UDC, style=style(doc, "Normal"))     # як у шаблоні: 14 пт, праворуч, 5 пт після
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(5)
    p.runs[0].font.size = Pt(14)
    doc.add_paragraph(TITLE, style=S["h1"])
    doc.add_paragraph(AUTHORS, style=S["auth"])
    doc.add_paragraph(ORG, style=S["org"])
    doc.add_paragraph("", style=style(doc, "Normal"))  # порожній рядок перед текстом, як у шаблоні
    for s in INTRO + ANALYSIS + ARCH_TEXT + ARCH_TEXT2:
        doc.add_paragraph(s, style=S["text"])
    pic = doc.add_paragraph(style=S["cap"])                      # рисунок: по центру, 5 пт до
    pic.paragraph_format.space_before, pic.paragraph_format.space_after = Pt(5), Pt(5)
    pic.paragraph_format.keep_with_next = True
    pic.add_run().add_picture(str(fig_png), width=Cm(FIG_W))
    cap = doc.add_paragraph(CAPTION, style=S["cap"])              # підпис: 5 пт до, 10 пт після
    cap.paragraph_format.space_before, cap.paragraph_format.space_after = Pt(5), Pt(10)
    for s in CONCLUSION:
        doc.add_paragraph(s, style=S["text"])
    doc.add_paragraph("Література", style=S["h2"])
    for i, r in enumerate(REFS, 1):
        doc.add_paragraph(f"{i}. {r}", style=S["ref"])
    cp = doc.core_properties
    cp.author, cp.title, cp.last_modified_by = "Петро Кравченко", TITLE, "Петро Кравченко"
    return doc


def verify(docx_path, pdf_path):
    text = "\n".join(p.text for p in Document(docx_path).paragraphs)
    body = "\n".join(INTRO + ANALYSIS)
    cited = {int(n) for grp in re.findall(r"\[([\d,\s–-]+)\]", body) for n in re.split(r"[,\s–-]+", grp) if n}
    rest = "\n".join(ARCH_TEXT + ARCH_TEXT2 + CONCLUSION)
    errs = []
    if cited != set(range(1, len(REFS) + 1)):
        errs.append(f"посилання в абзацах вступу/аналізу {sorted(cited)} ≠ 1…{len(REFS)}")
    first = [int(m) for m in re.findall(r"\[(\d+)\]", body)]
    order = list(dict.fromkeys(first))
    if order != sorted(order):
        errs.append(f"нумерація джерел не за першим згадуванням: {order}")
    if re.search(r"\[\d", rest):
        errs.append("посилання на джерела поза абзацами 1–2 (вимога завдання)")
    if "Рис. 1" not in "\n".join(ARCH_TEXT):
        errs.append("немає посилання на Рис. 1 у тексті")
    if re.search(r"[ыэъё]", text, re.I):
        errs.append("російські літери в тексті")
    if re.search(r"  |\t", text):
        errs.append("подвійні пробіли або табуляція (заборонено шаблоном)")
    pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf_path], capture_output=True,
                                                             text=True).stdout).group(1))
    bbox = subprocess.run(["pdftotext", "-bbox", "-f", str(pages), "-l", str(pages), pdf_path, "-"],
                          capture_output=True, text=True).stdout
    page_h = float(re.search(r'<page width="[\d.]+" height="([\d.]+)"', bbox).group(1))
    fill = max(float(y) for y in re.findall(r'yMax="([\d.]+)"', bbox)) / page_h  # низ останнього рядка
    if pages not in (1, 2):
        errs.append(f"обсяг {pages} стор. — дозволено 1 або 2 повні сторінки")
    elif fill < 0.85:
        errs.append(f"остання сторінка заповнена на {fill:.0%} висоти (потрібна повна сторінка)")
    return pages, fill, errs


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    r = render(FIG, HERE / "01_architecture")
    pt = r["portrait"]["font_pt"] * FIG_W / r["portrait"]["width_cm"]  # шрифт рисунка на ширині FIG_W
    print(f"рисунок: {r['native_px']} px, шрифт на {FIG_W} см ≈ {pt:.1f} пт, {r['portrait']['dpi']} dpi")
    if pt < 10:
        raise SystemExit("рисунок нечитабельний (< 10 пт) — зменшіть colw/bw")
    doc = build(HERE / "01_architecture.png")
    docx = OUT / f"{BASE}.docx"
    doc.save(docx)
    # x2t не вміє записувати .doc (Word 97–2003), LibreOffice на машині немає → .docx; у .doc — «Зберегти як» у Word
    x2t(docx, OUT / f"{BASE}.pdf", 513)
    pages, nlines, errs = verify(docx, str(OUT / f"{BASE}.pdf"))
    print(f"{BASE}.docx: {pages} стор., остання заповнена до {nlines:.0%} висоти")
    for e in errs:
        print("FAIL:", e)
    sys.exit(1 if errs else 0)
